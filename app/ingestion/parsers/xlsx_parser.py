"""XLSX parser: workbook -> one Markdown table per sheet.

Each worksheet becomes one ``ParsedSection`` (``section_path`` = sheet name)
holding a Markdown table, so spreadsheets chunk/embed/cite like any other
document. Rules (skill, section 1):

* Ingest computed values, not formulas (``data_only=True`` reads the cached
  result Excel last saved; a formula cell never opened in Excel yields "").
* Forward-fill merged cells from their anchor so every row is self-contained.
* Oversized tables are split later by chunking.py's row-group logic, which
  repeats the header row in every group — no size handling needed here.
"""

from __future__ import annotations

import unicodedata
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from app.ingestion.parsers.docx_parser import title_from_path
from app.models.schemas import DocType, ParsedDocument, ParsedSection


def _format_cell(value: Any) -> str:
    """Render one cell value as Markdown-table-safe text."""
    if value is None:
        return ""
    # Excel stores many integers as floats; drop the pointless ".0".
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        # Pure dates come back as midnight datetimes; show just the date.
        value = (
            value.date().isoformat() if value.time() == time(0) else value.isoformat()
        )
    elif isinstance(value, (date, time)):
        value = value.isoformat()
    text = unicodedata.normalize("NFC", str(value))
    # Keep the table grid intact: cells must stay single-line and pipe-free.
    return " ".join(text.split()).replace("|", "\\|")


def _sheet_rows(ws: Any) -> list[list[str]]:
    """Read a worksheet into formatted rows, forward-filling merged cells."""
    rows: list[list[Any]] = [[cell.value for cell in row] for row in ws.iter_rows()]
    for rng in ws.merged_cells.ranges:
        anchor = rows[rng.min_row - 1][rng.min_col - 1]
        for r in range(rng.min_row - 1, rng.max_row):
            for c in range(rng.min_col - 1, rng.max_col):
                rows[r][c] = anchor
    formatted = [[_format_cell(v) for v in row] for row in rows]
    # Drop fully-empty rows and columns (spreadsheets often have stray blanks).
    formatted = [row for row in formatted if any(row)]
    if not formatted:
        return []
    keep = [
        c
        for c in range(max(len(r) for r in formatted))
        if any(c < len(row) and row[c] for row in formatted)
    ]
    return [[row[c] if c < len(row) else "" for c in keep] for row in formatted]


def sheet_to_markdown(ws: Any) -> str:
    """Render one worksheet as a Markdown table (first row = header).

    Returns "" for an empty sheet.
    """
    rows = _sheet_rows(ws)
    if not rows:
        return ""
    header, *body = rows
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    lines += ["| " + " | ".join(row) + " |" for row in body]
    return "\n".join(lines)


def parse_xlsx(
    path: str | Path, *, doc_type: DocType = DocType.SPREADSHEET
) -> ParsedDocument:
    """Parse a ``.xlsx`` workbook into one Markdown-table section per sheet."""
    from openpyxl import load_workbook  # imported lazily; heavy-ish dependency

    p = Path(path)
    wb = load_workbook(p, data_only=True)
    sections = [
        ParsedSection(section_path=unicodedata.normalize("NFC", ws.title), text=table)
        for ws in wb.worksheets
        if (table := sheet_to_markdown(ws))
    ]
    # Prefer the workbook's Title property: filenames are usually ASCII-only,
    # so it is the only place a proper diacritics title (used in citations) survives.
    title = (wb.properties.title or "").strip() or title_from_path(p)
    return ParsedDocument(
        doc_title=unicodedata.normalize("NFC", title),
        doc_type=doc_type,
        sections=sections,
        source_path=str(p),
    )
